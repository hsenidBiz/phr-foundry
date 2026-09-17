#!/usr/bin/env python
"""Build an organization-standard PeoplesHR Configuration Document (.docx)
from a standardized configuration-document Markdown file.

Usage:
    python build_docx.py <document.md> [-o <output.docx>]

The Markdown must start with a front-matter block (see assets/skeleton.md).
Cover page, version/copyright page, table of contents, heading numbering,
header/footer and "Page X of Y" are generated here - never write them in the .md.

Supported Markdown subset:
    # / ## / ###            unnumbered headings (numbers are added here)
    # Change Control        always last, never numbered
    # Appendix A - Title    appendix heading, never numbered
    paragraphs with **bold**, *italic*, `code`
    - bullet / 1. numbered  (indent 2 spaces per nesting level)
    | table |               GitHub pipe tables (first row = header)
    ```sql ... ```          code blocks
    > **Note:** text        callouts: Note / Important / Warning / Recommendation
    ![Figure 1: Caption.](images/file.png)   figure + caption below
    <!-- pagebreak -->      explicit page break
Requires: pip install python-docx
"""
import argparse
import datetime
import json
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.section import WD_SECTION  # noqa: F401
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    sys.exit("python-docx is required: python -m pip install --user python-docx")

SKILL_DIR = Path(__file__).resolve().parent.parent
HOUSE = json.loads((SKILL_DIR / "assets" / "house.json").read_text(encoding="utf-8"))
LOGO = SKILL_DIR / "assets" / "peopleshr-logo.png"
BRAND = HOUSE["colors"]["brand"]
REQUIRED_META = ["feature_id", "title", "module", "release", "doc_version"]


# ---------------------------------------------------------------- low-level XML helpers
def set_fonts(rpr_owner, name):
    """Force a font on a style or run, removing theme-font overrides."""
    rpr = rpr_owner.get_or_add_rPr() if hasattr(rpr_owner, "get_or_add_rPr") else rpr_owner
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def set_cell_borders(cell, **edges):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge, (size, color) in edges.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single" if size else "nil")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)
        borders.append(el)
    tcpr.append(borders)


def set_table_borders(table, color, size=4):
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)
        borders.append(el)
    tblpr.append(borders)


def para_border(paragraph, edge, color=BRAND, size=8):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), "1")
    el.set(qn("w:color"), color)
    pbdr.append(el)
    ppr.append(pbdr)


def add_field(paragraph, instr, placeholder="", font_size=None, color=None):
    """Insert a complex field (PAGE, NUMPAGES, TOC, SEQ ...)."""
    def run_with(child):
        r = paragraph.add_run()
        if font_size:
            r.font.size = font_size
        if color:
            r.font.color.rgb = color
        r._r.append(child)
        return r

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    run_with(begin)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    run_with(it)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run_with(sep)
    r = paragraph.add_run(placeholder)
    if font_size:
        r.font.size = font_size
    if color:
        r.font.color.rgb = color
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_with(end)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def no_split(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


# ---------------------------------------------------------------- inline markdown
INLINE = re.compile(r"(`[^`]+`|\*\*.+?\*\*|(?<![\*\w])\*[^*\s][^*]*?\*(?!\*))")


def add_inline(paragraph, text, bold=False, italic=False, size=None, color=None):
    text = re.sub(r"<(https?://[^>\s]+)>", r"\1", text)  # markdown autolink -> plain URL
    for part in INLINE.split(text):
        if not part:
            continue
        b, i, code = bold, italic, False
        if part.startswith("`") and part.endswith("`") and len(part) > 1:
            part, code = part[1:-1], True
        elif part.startswith("**") and part.endswith("**") and len(part) > 4:
            add_inline(paragraph, part[2:-2], True, italic, size, color)
            continue
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            add_inline(paragraph, part[1:-1], bold, True, size, color)
            continue
        run = paragraph.add_run(part.replace("\\|", "|"))
        run.bold, run.italic = b or None, i or None
        if size:
            run.font.size = size
        if color:
            run.font.color.rgb = color
        if code:
            set_fonts(run._r, HOUSE["fonts"]["code"])
            run.font.size = Pt(9.5) if not size else size
            run.font.color.rgb = RGBColor.from_string("1F3864")


# ---------------------------------------------------------------- document setup
def setup_styles(doc):
    body, heading = HOUSE["fonts"]["body"], HOUSE["fonts"]["heading"]
    normal = doc.styles["Normal"]
    set_fonts(normal.element, body)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for level, size in ((1, 14), (2, 13), (3, 11)):
        st = doc.styles[f"Heading {level}"]
        set_fonts(st.element, heading)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = RGBColor.from_string(BRAND)
        st.paragraph_format.space_before = Pt(18 if level == 1 else 12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
    cap = doc.styles["Caption"]
    set_fonts(cap.element, body)
    cap.font.size = Pt(9)
    cap.font.italic = True
    cap.font.bold = False
    cap.font.color.rgb = RGBColor.from_string("404040")
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)


def setup_page(doc, meta, base_name):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))
    sec.different_first_page_header_footer = True
    blue = RGBColor.from_string(BRAND)

    # first page (cover): classification only
    fp = sec.first_page_footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(HOUSE["classification"])
    r.font.size, r.font.color.rgb = Pt(9), blue

    # other pages: header rule + running title, footer rule + file name + Page X of Y
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run(f"Feature ID – {meta['feature_id']}  |  {meta['module']}")
    r.font.size, r.font.color.rgb = Pt(9), blue
    para_border(hp, "bottom")

    fp = sec.footer.paragraphs[0]
    fp.style = doc.styles["Normal"]  # drop the Footer style's centre tab
    fp.paragraph_format.space_after = Pt(0)
    para_border(fp, "top")
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    r = fp.add_run(base_name)
    r.font.size, r.font.color.rgb = Pt(9), blue
    r = fp.add_run("\tPage ")
    r.font.size, r.font.color.rgb = Pt(10), blue
    add_field(fp, "PAGE", "1", Pt(10), blue)
    r = fp.add_run(" of ")
    r.font.size, r.font.color.rgb = Pt(10), blue
    add_field(fp, "NUMPAGES", "1", Pt(10), blue)
    cp = sec.footer.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cp.add_run(HOUSE["classification"])
    r.font.size, r.font.color.rgb = Pt(8), blue

    # ask Word to refresh TOC / page fields when the file is opened
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build_cover(doc, meta, figure_count):
    blue = RGBColor.from_string(BRAND)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if LOGO.exists():
        p.add_run().add_picture(str(LOGO), width=Inches(2.74))
    lines = [f"Feature ID – {meta['feature_id']}", meta["title"], meta["module"]]
    for idx, text in enumerate(lines):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if idx == 0:
            p.paragraph_format.space_before = Pt(200)
        r = p.add_run(text)
        r.bold, r.italic = True, True
        r.font.size, r.font.color.rgb = Pt(16), blue
    rel = meta.get("release", "")
    if meta.get("release_version"):
        rel = f"{rel} (Version {meta['release_version']})"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(f"Release {rel}")
    r.italic, r.font.size, r.font.color.rgb = True, Pt(12), blue
    page_break(doc)

    # version + copyright + TOC page
    p = doc.add_paragraph()
    r = p.add_run(f"The version of this document is {meta['doc_version']}")
    r.bold = True
    year = meta.get("year") or str(datetime.date.today().year)
    # one paragraph per line of the copyright text
    for line in HOUSE["copyright"].replace("{year}", year).split("\n"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(line)
        r.font.size = Pt(9)
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("Table of Contents")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(14), blue
    set_fonts(r._r, HOUSE["fonts"]["heading"])
    add_field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u',
              "Right-click here and choose Update Field to build the table of contents.")
    if figure_count >= 2:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        r = p.add_run("List of Figures")
        r.bold, r.font.size, r.font.color.rgb = True, Pt(14), blue
        set_fonts(r._r, HOUSE["fonts"]["heading"])
        add_field(doc.add_paragraph(), 'TOC \\h \\z \\c "Figure"',
                  "Right-click here and choose Update Field to build the list of figures.")
    page_break(doc)


# ---------------------------------------------------------------- block builders
def build_table(doc, rows):
    cols = max(len(r) for r in rows)
    rows = [r + [""] * (cols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, HOUSE["colors"]["table_border"])
    label_headers = {"item", "term", "attribute", "component", "step", "check", "diagnostic field", "gate"}
    label_col = cols == 2 and rows[0][0].strip("* ").lower() in label_headers
    two_col = label_col
    cell_size = Pt(9) if cols >= 5 else Pt(10)
    for ri, row in enumerate(rows):
        no_split(table.rows[ri])
        for ci, text in enumerate(row):
            cell = table.cell(ri, ci)
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            chunks = [c for c in re.split(r"<br\s*/?>", text.strip())]
            for k, chunk in enumerate(chunks):
                para = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
                para.paragraph_format.space_after = Pt(2)
                add_inline(para, chunk.strip(), bold=(ri == 0 or (two_col and ci == 0)), size=cell_size)
            if ri == 0:
                shade(cell, HOUSE["colors"]["table_header"])
        if ri == 0:
            repeat_header(table.rows[0])
    # column widths: every column fits its longest unbreakable word, the rest is
    # shared by typical text length (avoids "Te/st Ca/se" and "EIM_QUAL_ATTA/CHMENT")
    page = 6.5
    char_in = 0.068 if cols >= 5 else 0.075
    mins, avgs = [], []
    for ci in range(cols):
        words = []
        for r in rows:
            words += [(w, "`" in r[ci]) for w in re.sub(r"[*]|<br\s*/?>", " ", r[ci]).split()]
        need = max((len(w.strip("`")) * (char_in * 1.12 if code else char_in) for w, code in words), default=0.5)
        avg_len = sum(len(r[ci]) for r in rows) / len(rows)
        prose_floor = 1.4 if avg_len > 30 and "`" not in "".join(r[ci] for r in rows[1:]) else 0.6
        mins.append(max(need + 0.2, prose_floor))
        avgs.append(max(avg_len, 4))
    if label_col:
        mins, avgs = [max(mins[0], 1.9), mins[1]], [1, 3]
    if sum(mins) >= page:
        widths = [page * m / sum(mins) for m in mins]
    else:
        spare = page - sum(mins)
        widths = [m + spare * a / sum(avgs) for m, a in zip(mins, avgs)]
    table.autofit = False
    for row in table.rows:
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_code(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, HOUSE["colors"]["code_bg"])
    set_table_borders(table, "D9D9D9")
    first = True
    for line in lines or [""]:
        para = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        para.paragraph_format.space_after = Pt(0)
        run = para.add_run(line)
        set_fonts(run._r, HOUSE["fonts"]["code"])
        run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_callout(doc, lines):
    text = " ".join(l.strip() for l in lines)
    m = re.match(r"\*\*(Note|Important|Warning|Recommendation):\*\*\s*(.*)", text)
    kind, body = (m.group(1), m.group(2)) if m else ("Note", text)
    style = HOUSE["callouts"][kind]
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, style["fill"])
    none = (0, "auto")
    set_cell_borders(cell, left=(24, style["bar"]), top=none, bottom=none, right=none)
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(2)
    r = para.add_run(f"{kind}: ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(style["bar"])
    add_inline(para, body)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_figure(doc, caption, src, md_dir):
    path = (md_dir / src).resolve() if not Path(src).is_absolute() else Path(src)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    if path.exists():
        p.add_run().add_picture(str(path), width=Inches(6.0))
    else:
        r = p.add_run(f"[Missing image: {src}]")
        r.bold, r.font.color.rgb = True, RGBColor.from_string("C00000")
        print(f"WARNING: image not found: {path}", file=sys.stderr)
    m = re.match(r"Figure\s+(\d+)\s*[:.\-–]\s*(.*)", caption)
    num, text = (m.group(1), m.group(2)) if m else ("", caption)
    cap = doc.add_paragraph(style="Caption")
    cap.add_run("Figure ")
    add_field(cap, "SEQ Figure \\* ARABIC", num or "1")
    cap.add_run(f": {text}")


# ---------------------------------------------------------------- main conversion
def parse_front_matter(text):
    m = re.match(r"^---\s*\n(.*?)\n---\s*\n", text, re.S)
    if not m:
        sys.exit("ERROR: front matter block (--- ... ---) is missing at the top of the .md")
    meta = {}
    for line in m.group(1).splitlines():
        if ":" in line:
            k, v = line.split(":", 1)
            meta[k.strip()] = v.strip().strip('"').strip("'")
    missing = [k for k in REQUIRED_META if not meta.get(k)]
    if missing:
        sys.exit(f"ERROR: front matter missing: {', '.join(missing)}")
    return meta, text[m.end():]


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    return [c.strip() for c in re.split(r"(?<!\\)\|", line)]


def convert(md_path, out_path):
    md_path = Path(md_path).resolve()
    text = md_path.read_text(encoding="utf-8")
    meta, body = parse_front_matter(text)
    base_name = md_path.stem
    lines = body.splitlines()
    figure_count = sum(1 for l in lines if re.match(r"^\s*!\[", l))

    doc = Document()
    setup_styles(doc)
    setup_page(doc, meta, base_name)
    build_cover(doc, meta, figure_count)
    cp = doc.core_properties
    cp.title, cp.subject = meta["title"], meta["module"]
    cp.keywords = f"{meta['release']};{meta['feature_id']};Configuration Document"

    counters = [0, 0, 0]
    prefix, sub = None, None  # appendix letter and its sub-counters
    i, first_h1 = 0, True
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped or re.match(r"^<!--(?!\s*pagebreak).*-->$", stripped):
            i += 1
            continue
        if re.match(r"^<!--\s*pagebreak\s*-->$", stripped):
            page_break(doc)
            i += 1
            continue

        hm = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if hm:
            level = len(hm.group(1))
            title = re.sub(r"^\d+(\.\d+)*\.?\s+", "", hm.group(2)).strip()
            special = level == 1 and re.match(r"^(Change Control|Appendix\b)", title, re.I)
            if level == 1 and not first_h1 and re.match(r"^Change Control", title, re.I):
                page_break(doc)
            first_h1 = first_h1 and level != 1
            if special:
                # "Appendix A" -> sub-headings A.1, A.1.1; "Change Control" -> sub-headings unnumbered
                label = title
                am = re.match(r"^Appendix\s+([A-Z])\b", title, re.I)
                prefix = am.group(1).upper() if am else ""
                sub = [0, 0]
            elif level == 1:
                prefix, sub = None, None
                counters[0] += 1
                counters[1] = counters[2] = 0
                label = f"{counters[0]}  {title}"
            elif sub is not None:
                sub[level - 2] += 1
                if level == 2:
                    sub[1] = 0
                nums = ".".join(str(c) for c in sub[: level - 1])
                label = f"{prefix}.{nums}  {title}" if prefix else title
            else:
                counters[level - 1] += 1
                for k in range(level, 3):
                    counters[k] = 0
                label = ".".join(str(c) for c in counters[:level]) + "  " + title
            doc.add_heading(label, level=level)
            i += 1
            continue

        if stripped.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i].rstrip())
                i += 1
            build_code(doc, block)
            i += 1
            continue

        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\|?\s*:?-{3,}", lines[i].strip()):
                    rows.append(split_row(lines[i]))
                i += 1
            build_table(doc, rows)
            continue

        if stripped.startswith(">"):
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i].strip()[1:].strip())
                i += 1
            build_callout(doc, block)
            continue

        fm = re.match(r"^!\[(.*?)\]\((.*?)\)\s*$", stripped)
        if fm:
            build_figure(doc, fm.group(1), fm.group(2), md_path.parent)
            i += 1
            continue

        lm = re.match(r"^(\s*)([-*]|\d+[.)])\s+(.*)$", line)
        if lm:
            indent = len(lm.group(1).replace("\t", "  ")) // 2
            para = doc.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Inches(0.35 + 0.3 * indent)
            pf.first_line_indent = Inches(-0.25)
            pf.space_after = Pt(3)
            pf.tab_stops.add_tab_stop(Inches(0.35 + 0.3 * indent))
            marker = "•" if lm.group(2) in "-*" else lm.group(2).rstrip(")").rstrip(".") + "."
            para.add_run(marker + "\t")
            add_inline(para, lm.group(3))
            i += 1
            continue

        # paragraph: join soft-wrapped lines
        buf = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
                r"^(#{1,3}\s|```|\||>|!\[|\s*([-*]|\d+[.)])\s|<!--)", lines[i]):
            buf.append(lines[i].strip())
            i += 1
        para = doc.add_paragraph()
        add_inline(para, " ".join(buf))

    out_path = Path(out_path) if out_path else md_path.with_suffix(".docx")
    doc.save(str(out_path))
    print(f"Built {out_path}")
    if not HOUSE.get("copyright_confirmed"):
        print("NOTE: assets/house.json copyright text is not yet confirmed against the official wording.")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("markdown")
    ap.add_argument("-o", "--output")
    args = ap.parse_args()
    convert(args.markdown, args.output)
